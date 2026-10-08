# job_automation/src/job_automation/documents/docx_letter_renderer.py

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from config.letter_profile import (
    LETTER_PROFILE,
    LetterProfile,
)
from domain.models import (
    LetterContent,
    LetterContext,
)


FONT_NAME = "Times New Roman"


class DocxLetterRenderer:
    def __init__(
        self,
        profile: LetterProfile = LETTER_PROFILE,
    ) -> None:
        self.profile = profile

    def render(
        self,
        context: LetterContext,
        content: LetterContent,
        destination: Path,
    ) -> None:
        destination.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document = Document()

        self._configure_document(
            document
        )

        self._add_sender(
            document
        )

        self._add_recipient(
            document=document,
            context=context,
        )

        self._add_date(
            document=document,
            context=context,
        )

        self._add_subject(
            document=document,
            content=content,
        )

        self._add_greeting(
            document=document,
            content=content,
        )

        self._add_body(
            document=document,
            content=content,
        )

        self._add_closing(
            document=document,
            content=content,
        )

        self._add_signature(
            document=document,
            content=content,
        )

        document.save(
            destination
        )

    def _configure_document(
        self,
        document: Document,
    ) -> None:
        section = document.sections[0]

        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.3)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

        style = document.styles["Normal"]
        style.font.name = FONT_NAME
        style.font.size = Pt(10.5)

        style.element.rPr.rFonts.set(
            qn("w:eastAsia"),
            FONT_NAME,
        )

        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.08

    def _add_sender(
        self,
        document: Document,
    ) -> None:
        sender_lines = (
            self.profile.full_name,
            *self.profile.address_lines,
            self.profile.phone,
            self.profile.email,
            self.profile.portfolio,
            self.profile.github,
        )

        for line in filter(None, sender_lines):
            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = (
                Pt(0)
            )

            run = paragraph.add_run(
                line
            )

            self._configure_run(
                run
            )

    def _add_recipient(
        self,
        document: Document,
        context: LetterContext,
    ) -> None:
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)

        for line in context.recipient.postal_lines:
            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
            )

            paragraph.paragraph_format.space_after = (
                Pt(0)
            )

            run = paragraph.add_run(
                line
            )

            self._configure_run(run)

    def _add_date(
        self,
        document: Document,
        context: LetterContext,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(8)

        run = paragraph.add_run(
            f"{self.profile.city}, le {context.date_label}"
        )

        self._configure_run(run)

        attention = document.add_paragraph()

        attention.paragraph_format.space_after = Pt(8)

        run = attention.add_run(
            context.recipient.attention_line
        )

        self._configure_run(run)

    def _add_subject(
        self,
        document: Document,
        content: LetterContent,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(10)

        run = paragraph.add_run(
            f"Objet : {content.subject}"
        )

        run.bold = True

        self._configure_run(run)

    def _add_greeting(
        self,
        document: Document,
        content: LetterContent,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(8)

        run = paragraph.add_run(
            content.greeting
        )

        self._configure_run(run)

    def _add_body(
        self,
        document: Document,
        content: LetterContent,
    ) -> None:
        for paragraph_text in content.paragraphs:
            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
            )

            paragraph.paragraph_format.space_after = (
                Pt(7)
            )

            run = paragraph.add_run(
                paragraph_text
            )

            self._configure_run(run)

    def _add_closing(
        self,
        document: Document,
        content: LetterContent,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(14)

        run = paragraph.add_run(
            content.closing
        )

        self._configure_run(run)

    def _add_signature(
        self,
        document: Document,
        content: LetterContent,
    ) -> None:
        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            content.signature
        )

        self._configure_run(run)

    @staticmethod
    def _configure_run(run) -> None:
        run.font.name = FONT_NAME

        run_properties = (
            run.element.get_or_add_rPr()
        )

        run_fonts = run_properties.find(
            qn("w:rFonts")
        )

        if run_fonts is None:
            run_fonts = (
                run_properties.makeelement(
                    qn("w:rFonts"),
                    {},
                )
            )

            run_properties.append(
                run_fonts
            )

        run_fonts.set(
            qn("w:eastAsia"),
            FONT_NAME,
        )