# src/job_automation/documents/docx_cv_renderer.py

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from config.candidate_profile import (
    CANDIDATE_PROFILE,
    DEFAULT_SKILLS,
    CandidateProfile,
)
from documents.html_cv_renderer import experience_metadata
from domain.models import CvContent


FONT_NAME = "Times New Roman"
BLUE = RGBColor(0x1F, 0x3A, 0x5F)


class DocxCvRenderer:
    def __init__(
        self,
        profile: CandidateProfile = CANDIDATE_PROFILE,
    ) -> None:
        self.profile = profile

    def render(
        self,
        content: CvContent,
        destination: Path,
    ) -> None:
        destination.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document = Document()

        self._configure_document(
            document
        )

        self._add_header(
            document=document,
            subtitle=content.subtitle,
        )

        self._add_section_title(
            document,
            "FORMATION",
        )

        self._add_formations(
            document
        )

        if self.profile.references_line:
            self._add_section_title(
                document,
                "RÉFÉRENCES",
            )

            references_run = document.add_paragraph().add_run(
                self.profile.references_line
            )

            self._set_run_font(references_run)

        self._add_section_title(
            document,
            "COMPÉTENCES",
        )

        self._add_skills(
            document=document,
            skills=content.skills or DEFAULT_SKILLS,
        )

        self._add_section_title(
            document,
            "EXPÉRIENCES",
        )

        self._add_experiences(
            document=document,
            experiences=content.experiences,
        )

        document.save(
            destination
        )

    def _configure_document(
        self,
        document: Document,
    ) -> None:
        section = document.sections[0]

        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(0.7)
        section.left_margin = Cm(1.2)
        section.right_margin = Cm(1.2)

        normal_style = document.styles["Normal"]
        normal_style.font.name = FONT_NAME
        normal_style.font.size = Pt(9.5)

        normal_style.element.rPr.rFonts.set(
            qn("w:eastAsia"),
            FONT_NAME,
        )

        paragraph_format = (
            normal_style.paragraph_format
        )

        paragraph_format.space_after = Pt(1)
        paragraph_format.line_spacing = 1.0

    def _add_header(
        self,
        document: Document,
        subtitle: str,
    ) -> None:
        name_paragraph = document.add_paragraph()
        name_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        name_run = name_paragraph.add_run(
            self.profile.name
        )

        name_run.bold = True
        name_run.font.size = Pt(18)
        name_run.font.color.rgb = BLUE

        self._set_run_font(name_run)

        subtitle_paragraph = (
            document.add_paragraph()
        )

        subtitle_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        subtitle_run = subtitle_paragraph.add_run(
            subtitle
        )

        subtitle_run.bold = True
        subtitle_run.font.size = Pt(11)

        self._set_run_font(
            subtitle_run
        )

        contact_paragraph = (
            document.add_paragraph()
        )

        contact_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        contact_run = contact_paragraph.add_run(
            self.profile.contact
        )

        contact_run.font.size = Pt(9)

        self._set_run_font(
            contact_run
        )

    def _add_section_title(
        self,
        document: Document,
        title: str,
    ) -> None:
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(
            title
        )

        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = BLUE

        self._set_run_font(run)

    def _add_formations(
        self,
        document: Document,
    ) -> None:
        for formation in self.profile.formations:
            title_paragraph = (
                document.add_paragraph()
            )

            title_run = title_paragraph.add_run(
                formation.title
            )

            title_run.bold = True

            self._set_run_font(
                title_run
            )

            detail = formation.detail

            detail_paragraph = (
                document.add_paragraph()
            )

            detail_run = detail_paragraph.add_run(
                detail
            )

            detail_run.italic = True

            self._set_run_font(
                detail_run
            )

    def _add_skills(
        self,
        document: Document,
        skills: dict[str, list[str]],
    ) -> None:
        for category, bullets in skills.items():
            if not bullets:
                continue

            title_paragraph = document.add_paragraph()
            title_paragraph.paragraph_format.space_before = Pt(2)

            title_run = title_paragraph.add_run(category)
            title_run.bold = True
            title_run.font.color.rgb = BLUE

            self._set_run_font(title_run)

            for bullet in bullets:
                self._add_markdown_bullet(
                    document=document,
                    text=str(bullet),
                )

    def _add_experiences(
        self,
        document: Document,
        experiences: list[dict],
    ) -> None:
        for experience in experiences:
            title_paragraph = (
                document.add_paragraph()
            )

            title_paragraph.paragraph_format.space_before = (
                Pt(3)
            )

            title_run = title_paragraph.add_run(
                str(
                    experience.get(
                        "titre",
                        "",
                    )
                )
            )

            title_run.bold = True

            self._set_run_font(
                title_run
            )

            metadata = experience_metadata(experience)

            if metadata:
                metadata_paragraph = (
                    document.add_paragraph()
                )

                metadata_run = (
                    metadata_paragraph.add_run(
                        metadata
                    )
                )

                metadata_run.italic = True

                self._set_run_font(
                    metadata_run
                )

            for bullet in experience.get(
                "bullets",
                [],
            ):
                self._add_markdown_bullet(
                    document=document,
                    text=str(bullet),
                )

    def _add_markdown_bullet(
        self,
        document: Document,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.paragraph_format.space_after = Pt(0)

        parts = re.split(
            r"(\*\*.*?\*\*)",
            text,
        )

        for part in parts:
            if not part:
                continue

            is_bold = (
                part.startswith("**")
                and part.endswith("**")
            )

            displayed_text = (
                part[2:-2]
                if is_bold
                else part
            )

            run = paragraph.add_run(
                displayed_text
            )

            run.bold = is_bold

            self._set_run_font(run)

    @staticmethod
    def _set_run_font(run) -> None:
        run.font.name = FONT_NAME

        run_properties = (
            run.element.get_or_add_rPr()
        )

        run_fonts = run_properties.find(
            qn("w:rFonts")
        )

        if run_fonts is None:
            run_fonts = (
                run_properties.makeelement(
                    qn("w:rFonts"),
                    {},
                )
            )

            run_properties.append(
                run_fonts
            )

        run_fonts.set(
            qn("w:eastAsia"),
            FONT_NAME,
        )
